import csv
import io
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.training import TrainingPlan, TrainingWeek, User, Workout, WorkoutStatus


class PlanImportError(ValueError):
    pass


class TrainingPlanImportService:
    """Import CSV/XLSX training plans into normalized plan, week, and workout rows."""

    def import_plan_file(
        self,
        db: Session,
        user: User,
        file_bytes: bytes,
        filename: str,
        replace_existing: bool = False,
    ) -> TrainingPlan:
        if replace_existing:
            for plan in list(user.training_plans):
                db.delete(plan)
            db.flush()

        rows = self._read_rows(file_bytes, filename)
        if not rows:
            raise PlanImportError("The file did not contain any workout rows.")

        default_start = self._extract_plan_start(rows)
        rows = self._expand_weekly_summary_rows(rows)
        workouts = [
            self._parse_workout_row(row, index + 2, default_start)
            for index, row in enumerate(rows)
        ]
        scheduled_dates = [workout["scheduled_date"] for workout in workouts]
        starts_on = min(scheduled_dates) if scheduled_dates else date.today()
        plan_title = self._extract_plan_title(rows, filename)

        plan = TrainingPlan(
            user=user,
            title=plan_title,
            description=f"Imported from {filename}.",
            goal_race=self._first_value(rows, ["goal_race", "race", "goal"]),
            starts_on=starts_on,
        )

        weeks_by_number: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for workout in workouts:
            week_number = workout.pop("week_number") or self._week_number_from_date(
                starts_on,
                workout["scheduled_date"],
            )
            weeks_by_number[week_number].append(workout)

        for week_number in sorted(weeks_by_number):
            week_workouts = sorted(
                weeks_by_number[week_number],
                key=lambda workout: workout["scheduled_date"],
            )
            week = TrainingWeek(
                plan=plan,
                week_number=week_number,
                starts_on=min(workout["scheduled_date"] for workout in week_workouts),
                focus=self._week_focus(rows, week_number),
            )
            for workout in week_workouts:
                week.workouts.append(Workout(**workout))
            plan.weeks.append(week)

        db.add(plan)
        db.commit()
        db.refresh(plan)
        return plan

    def _read_rows(self, file_bytes: bytes, filename: str) -> list[dict[str, str]]:
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv":
            return self._read_csv(file_bytes)
        if suffix in {".xlsx", ".xlsm"}:
            return self._read_xlsx(file_bytes)
        if suffix == ".xls":
            raise PlanImportError("Legacy .xls files are not supported yet. Save as .xlsx or .csv.")
        raise PlanImportError("Unsupported file type. Upload a .csv or .xlsx file.")

    def _read_csv(self, file_bytes: bytes) -> list[dict[str, str]]:
        text = file_bytes.decode("utf-8-sig")
        sample = text[:2048]
        dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        if not reader.fieldnames:
            raise PlanImportError("CSV file is missing a header row.")
        return [self._normalize_row(row) for row in reader if any(row.values())]

    def _read_xlsx(self, file_bytes: bytes) -> list[dict[str, str]]:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise PlanImportError("Spreadsheet is empty.")

        header_index = self._find_header_index(rows)
        headers = [self._normalize_key(value) for value in rows[header_index]]
        parsed_rows = []
        for raw_row in rows[header_index + 1 :]:
            row = {
                headers[index]: self._cell_to_string(value)
                for index, value in enumerate(raw_row)
                if index < len(headers) and headers[index]
            }
            if any(row.values()):
                parsed_rows.append(row)
        return parsed_rows

    def _find_header_index(self, rows: list[tuple[Any, ...]]) -> int:
        for index, row in enumerate(rows[:10]):
            normalized = {self._normalize_key(value) for value in row}
            if normalized & {"date", "scheduled_date", "workout", "title"}:
                return index
        raise PlanImportError("Could not find a header row with date/title workout columns.")

    def _parse_workout_row(
        self,
        row: dict[str, str],
        row_number: int,
        default_start: date,
    ) -> dict[str, Any]:
        title = self._value(
            row,
            ["title", "workout", "name", "session", "run", "quality_session", "long_easy_run"],
        )
        scheduled_date = self._parse_scheduled_date(row, row_number, default_start)
        if not title:
            raise PlanImportError(f"Row {row_number} is missing a workout title.")

        return {
            "week_number": self._parse_int(self._value(row, ["week", "week_number"])),
            "title": title,
            "description": self._value(row, ["description", "details", "notes", "workout_notes"]),
            "planned_distance_km": self._parse_float(
                self._value(row, ["planned_distance_km", "distance_km", "distance", "km", "title"])
            ),
            "planned_duration_min": self._parse_int(
                self._value(row, ["planned_duration_min", "duration_min", "duration", "minutes"])
            ),
            "planned_pace": self._value(row, ["planned_pace", "pace", "target_pace"]),
            "planned_hr_zone": self._value(row, ["planned_hr_zone", "hr_zone", "zone"]),
            "scheduled_date": scheduled_date,
            "status": self._parse_status(self._value(row, ["status"])),
            "actual_distance_km": self._parse_float(
                self._value(row, ["actual_distance_km", "actual_distance"])
            ),
            "actual_duration_min": self._parse_int(
                self._value(row, ["actual_duration_min", "actual_duration"])
            ),
            "actual_avg_hr": self._parse_int(self._value(row, ["actual_avg_hr", "avg_hr"])),
            "actual_avg_cadence": self._parse_int(
                self._value(row, ["actual_avg_cadence", "avg_cadence"])
            ),
            "actual_avg_pace": self._value(row, ["actual_avg_pace", "avg_pace"]),
            "notes": self._value(row, ["notes", "comments"]),
        }

    def _normalize_row(self, row: dict[str, Any]) -> dict[str, str]:
        return {self._normalize_key(key): self._cell_to_string(value) for key, value in row.items()}

    def _normalize_key(self, value: Any) -> str:
        key = self._cell_to_string(value).lower()
        key = re.sub(r"[^a-z0-9]+", "_", key)
        return key.strip("_")

    def _cell_to_string(self, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return str(value).strip()

    def _value(self, row: dict[str, str], keys: list[str]) -> str | None:
        for key in keys:
            value = row.get(key)
            if value:
                return value
        return None

    def _first_value(self, rows: list[dict[str, str]], keys: list[str]) -> str | None:
        for row in rows:
            value = self._value(row, keys)
            if value:
                return value
        return None

    def _expand_weekly_summary_rows(self, rows: list[dict[str, str]]) -> list[dict[str, str]]:
        expanded_rows: list[dict[str, str]] = []
        for row in rows:
            has_date = bool(
                self._value(row, ["scheduled_date", "workout_date", "date", "fecha"])
                or self._value(row, ["day", "weekday", "dia", "día"])
            )
            has_weekly_summary = bool(
                self._value(row, ["quality_session"]) or self._value(row, ["long_easy_run"])
            )
            if has_date or not has_weekly_summary:
                expanded_rows.append(row)
                continue

            week = self._value(row, ["week", "week_number", "semana"]) or "1"
            phase = self._value(row, ["phase", "focus", "week_focus"])
            quality = self._value(row, ["quality_session"])
            long_run = self._value(row, ["long_easy_run"])

            if quality:
                expanded_rows.append(
                    {
                        **row,
                        "week": week,
                        "day": "Wednesday",
                        "title": "Quality Session",
                        "description": quality,
                        "focus": phase or "",
                    }
                )
            if long_run:
                expanded_rows.append(
                    {
                        **row,
                        "week": week,
                        "day": "Sunday",
                        "title": "Long/Easy Run",
                        "description": long_run,
                        "focus": phase or "",
                    }
                )

        return expanded_rows

    def _extract_plan_start(self, rows: list[dict[str, str]]) -> date:
        value = self._first_value(
            rows,
            ["plan_start", "plan_start_date", "start", "start_date", "starts_on"],
        )
        if value:
            try:
                return self._parse_date(value, 1)
            except PlanImportError:
                pass
        today = date.today()
        return today - timedelta(days=today.weekday())

    def _extract_plan_title(self, rows: list[dict[str, str]], filename: str) -> str:
        title = self._first_value(rows, ["plan_title", "plan", "training_plan"])
        if title:
            return title
        return Path(filename).stem.replace("_", " ").replace("-", " ").title()

    def _week_focus(self, rows: list[dict[str, str]], week_number: int) -> str | None:
        for row in rows:
            if self._parse_int(self._value(row, ["week", "week_number"])) == week_number:
                return self._value(row, ["focus", "week_focus"])
        return None

    def _parse_scheduled_date(
        self,
        row: dict[str, str],
        row_number: int,
        default_start: date,
    ) -> date:
        explicit_date = self._value(
            row,
            [
                "scheduled_date",
                "workout_date",
                "date",
                "fecha",
                "scheduled",
                "session_date",
            ],
        )
        if explicit_date:
            return self._parse_date(explicit_date, row_number)

        week_number = self._parse_int(self._value(row, ["week", "week_number", "semana"])) or 1
        day_value = self._value(row, ["day", "weekday", "dia", "día"])
        if day_value:
            day_delta = ((week_number - 1) * 7) + self._day_offset(day_value)
            return default_start + timedelta(days=day_delta)

        raise PlanImportError(
            f"Row {row_number} is missing a scheduled date. "
            "Add a date column, or provide week plus day/weekday columns."
        )

    def _parse_date(self, value: str | None, row_number: int) -> date:
        if not value:
            raise PlanImportError(f"Row {row_number} is missing a scheduled date.")
        for date_format in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y", "%d/%m/%y"):
            try:
                return datetime.strptime(value, date_format).date()
            except ValueError:
                continue
        raise PlanImportError(f"Row {row_number} has an unsupported date: {value}.")

    def _day_offset(self, value: str) -> int:
        normalized = self._normalize_key(value)
        day_offsets = {
            "monday": 0,
            "mon": 0,
            "lunes": 0,
            "tuesday": 1,
            "tue": 1,
            "martes": 1,
            "wednesday": 2,
            "wed": 2,
            "miercoles": 2,
            "miércoles": 2,
            "thursday": 3,
            "thu": 3,
            "jueves": 3,
            "friday": 4,
            "fri": 4,
            "viernes": 4,
            "saturday": 5,
            "sat": 5,
            "sabado": 5,
            "sábado": 5,
            "sunday": 6,
            "sun": 6,
            "domingo": 6,
        }
        if normalized in day_offsets:
            return day_offsets[normalized]

        parsed = self._parse_int(value)
        if parsed is not None:
            return max(min(parsed, 7), 1) - 1

        raise PlanImportError(f"Unsupported day value: {value}.")

    def _parse_float(self, value: str | None) -> float | None:
        if not value:
            return None
        match = re.search(r"\d+(?:\.\d+)?", value.replace(",", "."))
        return float(match.group()) if match else None

    def _parse_int(self, value: str | None) -> int | None:
        parsed = self._parse_float(value)
        return round(parsed) if parsed is not None else None

    def _parse_status(self, value: str | None) -> WorkoutStatus:
        normalized = (value or "planned").strip().lower()
        if normalized in {status.value for status in WorkoutStatus}:
            return WorkoutStatus(normalized)
        return WorkoutStatus.planned

    def _week_number_from_date(self, starts_on: date, scheduled_date: date) -> int:
        return max(((scheduled_date - starts_on).days // 7) + 1, 1)
